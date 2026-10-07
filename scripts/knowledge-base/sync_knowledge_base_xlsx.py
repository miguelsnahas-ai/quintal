#!/usr/bin/env python3
"""Regenerate the SQL to sync public.knowledge_chunks with a direct .xlsx
export of the source sheet ("base de conhecimento inicial para o RAG da IA
do Quintal").

Sibling of sync_knowledge_base.py (same table, same upsert philosophy —
sheet is the source of truth, re-run end to end, never hand-edit the
generated SQL), for the other input path: a .xlsx file (e.g. uploaded
directly in chat) instead of a Google Sheet's markdown export. Reading
the workbook directly is actually more robust for this source than the
markdown round-trip: no risk of a multi-line cell corrupting a markdown
table's row boundaries, and no need for the "spacer/alignment/header"
3-line-offset block parsing the other script does. Use this one whenever
you have the .xlsx itself; use sync_knowledge_base.py when you only have
MCP access to the Google Sheet (Google_Drive read_file_content).

HOW TO RUN A SYNC
  1. python3 scripts/knowledge-base/sync_knowledge_base_xlsx.py <file.xlsx> > /tmp/kb_sync.sql
  2. Skim the per-category row counts this script prints to stderr
     against the sheet's own "Leia-me" tab ("Contagem de linhas" row) —
     sanity check before applying anything.
  3. Apply /tmp/kb_sync.sql via the Supabase MCP tool `execute_sql`
     (project_id "izattwaiqjzhydzhxlns").
  4. knowledge_chunks_tsvector / search_knowledge_chunks need no change —
     the generated `search` column recomputes itself per upserted row.

WHAT IT DOES NOT TOUCH
  `status` (stays at its table default, 'published', for every synced
  row — same as sync_knowledge_base.py) and `image_url` (populated by
  scripts/knowledge-base/sync_activity_images.py from a different source,
  never by this script — deliberately absent from the INSERT/UPDATE
  column list so an existing image_url survives the upsert untouched).

ADDING A NEW TAB
  Add an entry to CONFIGS: (sheet tab name, category slug, title column
  name or None, age-min column name or None, age-max column name or
  None, tags column name or None). The ID column is always read as
  whichever column is first in that tab (position, not a literal "ID"
  string match) — some tabs (ex.: "9a BNCC - Objetivos EI") name it
  something else ("ID (código BNCC)"), and the BNCC objective codes
  themselves (ex.: EI02EO03) don't match a PREFIX-NUMBER shape, so this
  script accepts any non-empty first-column value as a valid ID rather
  than enforcing one fixed pattern — official codes must be preserved
  verbatim (the AI is expected to cite them, see the sheet's own
  "BNCC – como usar no Quintal" note).
"""

import sys
from collections import Counter

try:
    import openpyxl
except ImportError:  # pragma: no cover
    raise SystemExit("Missing dependency: pip install openpyxl")

# (sheet tab name, category slug, title column, age-min column, age-max column, tags column)
CONFIGS = [
    ("1 Materiais", "materiais", "Material", "Idade mín. (meses)", "Idade máx. (meses)", "Tags"),
    ("2 Músicas e brincadeiras", "brincadeiras", "Nome", "Idade mín. (meses)", "Idade máx. (meses)", "Tags"),
    ("3a Alimentação - alimentos", "alimentos", "Alimento", "Idade mín. (meses)", None, None),
    ("3b Alimentação - receitas", "receitas", "Receita", "Idade mín. (meses)", None, None),
    ("3c Métodos de introdução", "metodos_alimentacao", "Método", None, None, None),
    ("4 Janelas e rotinas de sono", "rotinas_sono", "Modelo de rotina", "Idade mín. (meses)", "Idade máx. (meses)", None),
    ("5 Formas de fazer dormir", "formas_de_dormir", "Método", "Idade mín. (meses)", "Idade máx. (meses)", None),
    ("6 Desenvolvimento", "desenvolvimento", "Área", "Idade mín. (meses)", "Idade máx. (meses)", None),
    ("7 Higiene e cuidado", "higiene", "Produto (tipo)", "Idade mín. (meses)", "Idade máx. (meses)", None),
    ("8 Passeios", "passeios", "Passeio", "Idade mín. (meses)", "Idade máx. (meses)", None),
    # Novas abas desta planilha — BNCC (referência curricular) e o
    # aprofundamento de Higiene (abas 10a–10i "aprofundam a aba 7",
    # nota da própria planilha — não a substituem).
    ("9a BNCC - Objetivos EI", "bncc_objetivos", None, "Idade mín. (meses)", "Idade máx. (meses)", None),
    ("9b BNCC - Fundamentos", "bncc_fundamentos", "Nome", None, None, None),
    ("10a Fraldas - troca", "higiene_fraldas_troca", "Título", "Idade mín. (meses)", "Idade máx. (meses)", "Tags"),
    ("10b Fraldas - tamanhos", "higiene_fraldas_tamanhos", "Item", "Idade mín. (meses)", "Idade máx. (meses)", None),
    ("10c Xixi e cocô", "higiene_xixi_coco", "Título", "Idade mín. (meses)", "Idade máx. (meses)", "Tags"),
    ("10d Higiene natural e desfralde", "higiene_natural_desfralde", "Título", "Idade mín. (meses)", "Idade máx. (meses)", "Tags"),
    ("10e Fralda de pano", "higiene_fralda_pano", "Título", "Idade mín. (meses)", "Idade máx. (meses)", "Tags"),
    ("10f Pele e área da fralda", "higiene_pele_fralda", "Quadro", "Idade mín. (meses)", "Idade máx. (meses)", "Tags"),
    ("10g Banho e cuidados diários", "higiene_banho", "Título", "Idade mín. (meses)", "Idade máx. (meses)", "Tags"),
    ("10h Produtos de higiene", "higiene_produtos", "Produto", "Idade mín. (meses)", "Idade máx. (meses)", None),
    ("10i Ingredientes", "higiene_ingredientes", "Ingrediente", None, None, None),
]


def clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        return str(value)
    value = str(value)
    if value in ("—", "-", "–", ""):
        return ""
    return value.strip()


def to_int(value) -> int | None:
    if isinstance(value, (int, float)):
        return int(value)
    text = clean(value)
    if not text:
        return None
    import re

    m = re.search(r"-?\d+", text)
    return int(m.group()) if m else None


def sql_str(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"


def sql_int_or_null(value) -> str:
    return str(value) if value is not None else "NULL"


def sql_text_array_or_null(tags):
    if not tags:
        return "NULL"
    return "ARRAY[" + ", ".join(sql_str(t) for t in tags) + "]::text[]"


def parse_sheet(path: str) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    records: list[dict] = []

    for tab_name, slug, title_col, age_min_col, age_max_col, tags_col in CONFIGS:
        if tab_name not in wb.sheetnames:
            raise SystemExit(f"Expected tab {tab_name!r} not found in workbook. Sheets present: {wb.sheetnames}")

        ws = wb[tab_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = [clean(h) for h in rows[0]]
        id_col_index = 0  # sempre a primeira coluna, qualquer que seja seu nome literal

        for raw_row in rows[1:]:
            rowmap = dict(zip(header, raw_row))
            rid = clean(raw_row[id_col_index]) if id_col_index < len(raw_row) else ""
            if not rid:
                continue

            title = clean(rowmap.get(title_col, "")) if title_col else ""
            title = title or rid
            age_min = to_int(rowmap.get(age_min_col)) if age_min_col else None
            age_max = to_int(rowmap.get(age_max_col)) if age_max_col else None

            tags = None
            if tags_col:
                raw_tags = clean(rowmap.get(tags_col, ""))
                if raw_tags:
                    tags = [t.strip() for t in raw_tags.split(";") if t.strip()]

            content_parts = []
            for h in header:
                if h == header[id_col_index]:
                    continue
                val = clean(rowmap.get(h, ""))
                if val:
                    content_parts.append(f"{h}: {val}")

            records.append(
                {
                    "id": rid,
                    "category": slug,
                    "title": title,
                    "age_min_months": age_min,
                    "age_max_months": age_max,
                    "tags": tags,
                    "content": "\n".join(content_parts),
                }
            )

    return records


def build_sql(records: list[dict]) -> str:
    values_sql = ",\n".join(
        "({id}, {category}, {title}, {age_min}, {age_max}, {tags}, {content})".format(
            id=sql_str(r["id"]),
            category=sql_str(r["category"]),
            title=sql_str(r["title"]),
            age_min=sql_int_or_null(r["age_min_months"]),
            age_max=sql_int_or_null(r["age_max_months"]),
            tags=sql_text_array_or_null(r["tags"]),
            content=sql_str(r["content"]),
        )
        for r in records
    )

    id_list_sql = ", ".join(sql_str(r["id"]) for r in records)

    return f"""-- Auto-generated by scripts/knowledge-base/sync_knowledge_base_xlsx.py —
-- do not hand-edit; re-run against a fresh .xlsx export instead.
-- Upserts every row currently in the sheet (insert new, update changed —
-- matched by id) and removes any row whose id is no longer present in
-- ANY tab of the sheet (i.e. deleted). status/image_url are left alone.

insert into public.knowledge_chunks
  (id, category, title, age_min_months, age_max_months, tags, content)
values
{values_sql}
on conflict (id) do update set
  category = excluded.category,
  title = excluded.title,
  age_min_months = excluded.age_min_months,
  age_max_months = excluded.age_max_months,
  tags = excluded.tags,
  content = excluded.content;

delete from public.knowledge_chunks
where id not in ({id_list_sql});
"""


def main() -> None:
    if len(sys.argv) != 2:
        print(f"Usage: {sys.argv[0]} <sheet.xlsx>", file=sys.stderr)
        raise SystemExit(2)

    records = parse_sheet(sys.argv[1])
    if not records:
        raise SystemExit("Parsed zero records — check the file and CONFIGS tab names match.")

    seen = Counter(r["id"] for r in records)
    dupes = [rid for rid, n in seen.items() if n > 1]
    if dupes:
        raise SystemExit(f"Duplicate IDs across tabs (would silently collide in the upsert): {dupes}")

    counts = Counter(r["category"] for r in records)
    print(f"Parsed {len(records)} rows across {len(counts)} categories:", file=sys.stderr)
    for slug, n in sorted(counts.items()):
        print(f"  {slug}: {n}", file=sys.stderr)

    print(build_sql(records))


if __name__ == "__main__":
    main()
